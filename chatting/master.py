"""The master spreadsheet: averages and ranks, written for the admin and read back for the video."""

from __future__ import annotations

import datetime
import shutil
from pathlib import Path
from typing import Callable, Optional

from . import log
from .config import VotingConfig
from .models import RankedSong, Results, Song, Vote
from .names import People
from .ratings import parse_rating_value
from .sources import SourceError
from .timestamps import format_time, parse_time

OWN = "own"
_FIXED = ["rank", "average", "votes", "song", "artist", "submitted by", "link", "start", "id"]

# song id -> voter -> (rating or None, comment)
VotesBySong = dict[int, dict[str, tuple[Optional[float], str]]]


def order_voters(
    voters: list[str], people: People, voting: VotingConfig, alphabetical: bool = True
) -> list[str]:
    """voter_order entries first (those present), then the rest A-Z (or as given)."""
    by_key = {people.key(v): v for v in voters}
    first: list[str] = []
    for name in voting.voter_order:
        v = by_key.get(people.key(name))
        if v is not None and v not in first:
            first.append(v)
    rest = [v for v in voters if v not in first]
    if alphabetical:
        rest.sort(key=lambda v: v.casefold())
    return first + rest


def rank_songs(
    songs: list[Song], voters: list[str], votes_by_song: VotesBySong, people: People,
    voting: VotingConfig, warn: Callable[[str], None] = log.warn, alphabetical: bool = True,
) -> Results:
    """Average (own ratings left out), rank (ties share one) and voter order."""
    ordered_voters = order_voters(voters, people, voting, alphabetical)
    ranked: list[RankedSong] = []
    for song in songs:
        cast = votes_by_song.get(song.id, {})
        ratings: dict[str, Optional[float]] = {}
        comments: dict[str, str] = {}
        excluded: set[str] = set()
        counted: list[float] = []
        for v in ordered_voters:
            rating, comment = cast.get(v, (None, ""))
            comments[v] = comment
            if voting.exclude_own and people.same(v, song.submitter):
                excluded.add(v)
                ratings[v] = None
                continue
            ratings[v] = rating
            if rating is not None:
                counted.append(rating)
        average = sum(counted) / len(counted) if counted else None
        if average is None:
            warn(f"Nobody rated song {song.id} ({song.label}).")
        ranked.append(RankedSong(song=song, ratings=ratings, comments=comments, excluded=excluded,
                                 average=average, votes=len(counted)))

    ranked.sort(key=lambda r: (r.average is None, -(r.average or 0.0), r.song.id))
    keys = [None if r.average is None else round(r.average, 6) for r in ranked]
    rated = sum(1 for k in keys if k is not None)
    for i, r in enumerate(ranked):
        if keys[i] is None:
            r.rank = rated + 1
        elif i > 0 and keys[i] == keys[i - 1]:
            r.rank = ranked[i - 1].rank
        else:
            r.rank = i + 1
    return Results(voters=ordered_voters, songs=ranked)


def compile_results(
    songs: list[Song], voters: list[str], votes: list[Vote], people: People,
    voting: VotingConfig, warn: Callable[[str], None] = log.warn,
) -> Results:
    by_song: VotesBySong = {}
    for v in votes:
        by_song.setdefault(v.song_id, {})[v.voter] = (v.rating, v.comment)
    return rank_songs(songs, voters, by_song, people, voting, warn)


# --- writing ---------------------------------------------------------------------------


def write_master(results: Results, path: Path, title: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        shutil.copy2(path, path.with_name(path.stem + ".backup.xlsx"))

    bold = Font(bold=True)
    green = PatternFill("solid", start_color="C6EFCE")
    red = PatternFill("solid", start_color="FFC7CE")
    grey = PatternFill("solid", start_color="D9D9D9")

    def header(ws, names, widths):
        ws.append(names)
        for c in ws[1]:
            c.font = bold
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    voters = results.voters
    header(ws, ["Rank", "Average", "Votes", "Song", "Artist", "Submitted by", "Link", "Start", "ID"] + voters,
           [6, 10, 7, 40, 28, 18, 40, 8, 6] + [12] * len(voters))
    for r in results.songs:
        s = r.song
        ws.append([r.rank, r.average, r.votes, s.title, s.artist, s.submitter, s.link,
                   format_time(s.start), s.id])
        row = ws.max_row
        ws.cell(row=row, column=2).number_format = "0.00"
        # Text, so a start time typed in Excel stays "1:23" instead of becoming 1 h 23 min.
        ws.cell(row=row, column=8).number_format = "@"
        if s.link.lower().startswith(("http://", "https://")):
            ws.cell(row=row, column=7).hyperlink = s.link
            ws.cell(row=row, column=7).font = Font(color="0563C1", underline="single")
        for j, v in enumerate(voters, 10):
            cell = ws.cell(row=row, column=j)
            if v in r.excluded:
                cell.value, cell.fill = OWN, grey
            else:
                rating = r.ratings.get(v)
                cell.value = rating
                if rating is not None and rating >= 10:
                    cell.fill = green
                elif rating is not None and rating < 5:
                    cell.fill = red
    ws.auto_filter.ref = ws.dimensions

    wc = wb.create_sheet("Comments")
    header(wc, ["Rank", "Song", "Artist", "Submitted by", "ID"] + voters, [6, 40, 28, 18, 6] + [40] * len(voters))
    for r in results.songs:
        s = r.song
        wc.append([r.rank, s.title, s.artist, s.submitter, s.id] + [r.comments.get(v, "") or None for v in voters])
        for j in range(6, 6 + len(voters)):
            wc.cell(row=wc.max_row, column=j).alignment = Alignment(wrap_text=True, vertical="top")

    wsub = wb.create_sheet("Submitters")
    header(wsub, ["Submitted by", "Songs", "Average", "Best rank", "Worst rank"], [20, 8, 10, 10, 10])
    groups: dict[str, list[RankedSong]] = {}
    for r in results.songs:
        groups.setdefault(r.song.submitter, []).append(r)
    rows = []
    for name, rs in groups.items():
        avgs = [r.average for r in rs if r.average is not None]
        rows.append((name, len(rs), sum(avgs) / len(avgs) if avgs else None,
                     min(r.rank for r in rs), max(r.rank for r in rs)))
    rows.sort(key=lambda t: (t[2] is None, -(t[2] or 0.0), t[0].casefold()))
    for row in rows:
        wsub.append(list(row))
        wsub.cell(row=wsub.max_row, column=3).number_format = "0.00"

    wv = wb.create_sheet("Voters")
    header(wv, ["Voter", "Ratings given", "Average given", "10s given", "Comments written"], [20, 14, 14, 10, 18])
    for v in voters:
        given = [r.ratings[v] for r in results.songs if r.ratings.get(v) is not None]
        wv.append([v, len(given), sum(given) / len(given) if given else None,
                   sum(1 for g in given if g >= 10), sum(1 for r in results.songs if r.comments.get(v))])
        wv.cell(row=wv.max_row, column=3).number_format = "0.00"

    if title:
        wb.properties.title = title
    wb.save(path)


# --- reading ---------------------------------------------------------------------------


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _start(value: object) -> Optional[float]:
    """Start time from a Start cell, including the shapes Excel turns typed times into."""
    if isinstance(value, datetime.datetime):
        value = value.time()
    if isinstance(value, datetime.timedelta):
        total = int(value.total_seconds())
        value = datetime.time((total // 3600) % 24, (total // 60) % 60, total % 60)
    if isinstance(value, datetime.time):
        # "1:23" typed into a time cell is stored as 01:23:00; nobody's clip starts an hour in.
        if value.second == 0:
            return value.hour * 60 + value.minute
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, float) and value != int(value):
        return parse_time(f"{value:.2f}")      # "1.30" typed into a number cell -> 1.3
    return parse_time(_text(value))


def read_master(
    path: Path, people: People, voting: VotingConfig, warn: Callable[[str], None] = log.warn
) -> Results:
    """Read the (possibly hand-edited) master back; averages and ranks are recomputed."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(Path(path), data_only=True)
    except Exception as exc:
        raise SourceError(f"Could not read the master sheet {path}: {exc}") from exc
    if "Ranking" not in wb.sheetnames:
        raise SourceError(f"{path} has no 'Ranking' sheet.")
    rows = list(wb["Ranking"].iter_rows(values_only=True))
    if not rows:
        raise SourceError(f"The 'Ranking' sheet in {path} is empty.")
    headers = [_text(h) for h in rows[0]]
    idx = {h.casefold(): i for i, h in enumerate(headers) if h and h.casefold() in _FIXED}
    if "id" not in idx or "song" not in idx:
        raise SourceError(f"The 'Ranking' sheet in {path} needs 'ID' and 'Song' columns.")
    voter_cols = [(i, h) for i, h in enumerate(headers) if h and h.casefold() not in _FIXED]
    voters = [h for _, h in voter_cols]

    def get(row: tuple, name: str) -> object:
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    songs: list[Song] = []
    votes_by_song: VotesBySong = {}
    for row in rows[1:]:
        sid_text = _text(get(row, "id"))
        if not sid_text:
            continue
        try:
            sid = int(float(sid_text))
        except ValueError:
            warn(f"Master sheet: ignoring row with unreadable ID '{sid_text}'.")
            continue
        song = Song(id=sid, title=_text(get(row, "song")), artist=_text(get(row, "artist")),
                    link=_text(get(row, "link")), submitter=people.canonical(_text(get(row, "submitted by"))),
                    start=_start(get(row, "start")))
        songs.append(song)
        cast = votes_by_song.setdefault(sid, {})
        for i, v in voter_cols:
            raw = _text(row[i]) if i < len(row) else ""
            if raw.casefold() == OWN:
                rating = None
            else:
                rating, problem = parse_rating_value(raw, voting.min, voting.max)
                if problem:
                    warn(f"Master sheet: {v}: '{song.label}': {problem}")
            cast[v] = (rating, "")

    if "Comments" in wb.sheetnames:
        crows = list(wb["Comments"].iter_rows(values_only=True))
        if crows:
            ch = [_text(h) for h in crows[0]]
            id_col = next((i for i, h in enumerate(ch) if h.casefold() == "id"), None)
            if id_col is not None:
                for row in crows[1:]:
                    try:
                        sid = int(float(_text(row[id_col])))
                    except (ValueError, IndexError):
                        continue
                    cast = votes_by_song.get(sid)
                    if cast is None:
                        continue
                    for i, h in enumerate(ch):
                        if h in cast and i < len(row):
                            cast[h] = (cast[h][0], _text(row[i]))

    return rank_songs(songs, voters, votes_by_song, people, voting, warn,
                      alphabetical=False)
