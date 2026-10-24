"""The anonymous rating sheet: writing the blank one, reading every filled-in copy."""

from __future__ import annotations

import csv
import difflib
import re
from pathlib import Path
from typing import Callable, Optional

import pandas as pd

from . import log
from .config import Config
from .links import is_url
from .models import Song, Vote
from .names import People, normalize
from .sources import SourceError, find_header_row, read_grid

_MISSING = {"", "-", "–", "—", "x", "n/a", "na", "forgot", "skip", "?", "nan"}
_RATING_HEADER = r"rating|grade|score|nota|notă|punctaj|rate|^note?$"
_SONG_OR_ID_HEADER = r"song|title|track|piesa|piesă|melodie|^#$|^id$|^nr?\.?$"
_SHEET_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm", ".ods"}


def _fmt(x: float) -> str:
    return f"{x:g}"


# --- writing ---------------------------------------------------------------------------


def write_rating_sheet(
    songs: list[Song], xlsx_path: Path, csv_path: Optional[Path] = None,
    vmin: float = 0, vmax: float = 10,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    ordered = sorted(songs, key=lambda s: s.id)
    headers = ["#", "Song", "Link", "Rating", "Comments"]
    xlsx_path = Path(xlsx_path)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ratings"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDE", (6, 50, 45, 10, 60)):
        ws.column_dimensions[col].width = width
    for i, song in enumerate(ordered, 2):
        ws.append([song.id, song.label, song.link, None, None])
        if is_url(song.link):
            link_cell = ws.cell(row=i, column=3)
            link_cell.hyperlink = song.link if song.link.lower().startswith("http") else "https://" + song.link.strip()
            link_cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=i, column=5).alignment = Alignment(wrap_text=True, vertical="top")
    if ordered:
        dv = DataValidation(
            type="decimal", operator="between", formula1=str(vmin), formula2=str(vmax),
            allow_blank=True, showInputMessage=True, showErrorMessage=True, errorStyle="warning",
            promptTitle="Rating", prompt=f"{_fmt(vmin)}–{_fmt(vmax)}, decimals allowed, e.g. 6.78",
            errorTitle="Rating", error=f"Expected a number from {_fmt(vmin)} to {_fmt(vmax)}.",
        )
        ws.add_data_validation(dv)
        dv.add(f"D2:D{len(ordered) + 1}")
    wb.save(xlsx_path)

    if csv_path:
        csv_path = Path(csv_path)
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            for song in ordered:
                w.writerow([song.id, song.label, song.link, "", ""])


# --- reading ---------------------------------------------------------------------------

_MONTH = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE_LIKE = [
    re.compile(r"^\d{4}-\d{1,2}-\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$"),
    re.compile(r"^\d{1,4}([-/.])\d{1,2}\1\d{1,4}$"),
    re.compile(r"^\d{1,2}[-/ ]" + _MONTH + r"\.?([-/ ]\d{2,4})?$", re.I),
    re.compile(r"^" + _MONTH + r"\.?[-/ ]\d{1,2}$", re.I),
]


def parse_rating_value(
    raw: object, vmin: float = 0, vmax: float = 10
) -> tuple[Optional[float], Optional[str]]:
    """(value, problem). Missing markers give (None, None); unusable text gives a problem."""
    text = "" if raw is None else str(raw).strip()
    if text.lower() in _MISSING:
        return None, None
    if any(rx.match(text) for rx in _DATE_LIKE):
        return None, f"'{text}' looks like a date (the spreadsheet turned it into one): retype it as a number"
    m = re.match(r"^([-+]?\d+(?:[.,]\d+)?)(.*)$", text, re.S)
    if not m:
        return None, f"can't read '{text}'"
    rest = m.group(2)
    den = re.match(r"^\s*/\s*(\d+(?:[.,]\d+)?)", rest)
    if den and float(den.group(1).replace(",", ".")) != vmax:
        return None, f"can't read '{text}'"
    if re.match(r"^\s*[-–—]\s*\d", rest) or re.match(r"^\d", rest):
        return None, f"can't read '{text}'"
    value = round(float(m.group(1).replace(",", ".")), 2)
    if value < vmin or value > vmax:
        return None, f"{_fmt(value)} is outside {_fmt(vmin)}–{_fmt(vmax)}"
    return value, None


def voter_name_from_path(path: Path) -> str:
    stem = Path(path).stem.strip()
    return re.sub(r"\s*\(\d+\)$", "", stem).strip()


def _find_cols(headers: list[str]) -> dict[str, Optional[int]]:
    cols: dict[str, Optional[int]] = {"rating": None, "id": None, "song": None, "comments": None}
    for i, h in enumerate(headers):
        if cols["rating"] is None and re.search(_RATING_HEADER, h, re.I):
            cols["rating"] = i
    for i, h in enumerate(headers):
        if i == cols["rating"]:
            continue
        if cols["id"] is None and re.search(r"^#$|^id$|^no\.?$|^nr\.?$|number", h, re.I):
            cols["id"] = i
        elif cols["comments"] is None and re.search(r"comment|notes|comentari|thoughts|review|remarks", h, re.I):
            cols["comments"] = i
        elif cols["song"] is None and re.search(r"song|title|track|piesa|piesă|melodie", h, re.I):
            cols["song"] = i
    return cols


def _matches(text: str, song: Song) -> bool:
    t = normalize(text)
    if not t:
        return False
    for cand in (normalize(song.label), normalize(song.title)):
        if cand and (t == cand or difflib.SequenceMatcher(None, t, cand).ratio() >= 0.8):
            return True
    return False


def parse_rating_sheet(
    grid: pd.DataFrame, voter: str, songs: list[Song], vmin: float = 0, vmax: float = 10,
    warn: Callable[[str], None] = log.warn,
) -> list[Vote]:
    by_id = {s.id: s for s in songs}
    ordered = sorted(songs, key=lambda s: s.id)
    hrow = find_header_row(grid, [_RATING_HEADER, _SONG_OR_ID_HEADER])
    if hrow is None:
        hrow = find_header_row(grid, [_RATING_HEADER])
    if hrow is None:
        warn(f"{voter}: no rating column found (looked for 'Rating', 'Nota', ...); ignoring this file.")
        return []
    headers = [str(c).strip() for c in grid.iloc[hrow].tolist()]
    cols = _find_cols(headers)
    rows = [
        [str(c).strip() for c in r]
        for r in grid.iloc[hrow + 1:].values.tolist()
        if any(str(c).strip() for c in r)
    ]

    def cell(r: list[str], key: str) -> str:
        i = cols[key]
        return r[i] if i is not None and i < len(r) else ""

    votes: list[Vote] = []
    seen: set[int] = set()
    warned_position = False
    for pos, r in enumerate(rows):
        id_text, text = cell(r, "id"), cell(r, "song")
        row_desc = text or " | ".join(c for c in r if c)
        sid: Optional[int] = None
        try:
            cand = int(float(id_text.replace(",", "."))) if id_text else None
        except ValueError:
            cand = None
        valid_id = cand if cand in by_id else None
        found = [s.id for s in ordered if _matches(text, s)] if text else []
        if valid_id is not None:
            if not text or valid_id in found:
                sid = valid_id
            elif len(found) == 1:
                sid = found[0]
                warn(f"{voter}: row #{valid_id} says '{text}', which is song {sid}; using song {sid}.")
            else:
                sid = valid_id
                warn(f"{voter}: row '{text}' does not match song #{valid_id} ({by_id[valid_id].label}); using the number.")
        elif len(found) == 1:
            sid = found[0]
        elif cols["id"] is None and len(rows) == len(songs):
            sid = ordered[pos].id
            if not warned_position:
                warned_position = True
                warn(f"{voter}: no number column; matching rows to songs by position.")
        else:
            warn(f"{voter}: could not tell which song this row is, skipped: '{row_desc}'")
            continue
        song = by_id[sid]
        if sid in seen:
            warn(f"{voter}: '{song.label}' appears twice; using the first.")
            continue
        seen.add(sid)
        raw = cell(r, "rating")
        value, problem = parse_rating_value(raw, vmin, vmax)
        if problem:
            warn(f"{voter}: '{song.label}': {problem}")
        comment = cell(r, "comments")
        if comment.lower() == "nan":
            comment = ""
        votes.append(Vote(voter=voter, song_id=sid, rating=value, raw=raw, comment=comment))
    return votes


def load_votes(
    cfg: Config, songs: list[Song], people: People, warn: Callable[[str], None] = log.warn
) -> tuple[list[str], list[Vote]]:
    sources: list[tuple[str, str]] = []   # (voter, file path or link)
    folder = cfg.path(cfg.input.ratings_dir)
    if folder.is_dir():
        for p in sorted(folder.iterdir(), key=lambda p: (voter_name_from_path(p).lower(), len(p.stem), p.name.lower())):
            if not p.is_file() or p.suffix.lower() not in _SHEET_SUFFIXES:
                continue
            if p.name.startswith((".", "~$")):
                continue
            if p.stem.strip().lower().replace("_", " ") == "rating sheet":
                continue
            sources.append((voter_name_from_path(p), str(p)))
    for name, link in cfg.input.rating_links.items():
        sources.append((name, link))
    if not sources:
        warn(f"No rating sheets found in {folder} (and no rating_links configured).")
        return [], []

    voters: list[str] = []
    keys: set[str] = set()
    votes: list[Vote] = []
    for raw_name, src in sources:
        voter = people.canonical(raw_name)
        if people.key(voter) in keys:
            warn(f"{voter}: more than one rating sheet; using the first one and ignoring {src}.")
            continue
        try:
            grid = read_grid(src, cfg.base_dir)
        except SourceError as exc:
            warn(f"{voter}: {exc}")
            continue
        keys.add(people.key(voter))
        voters.append(voter)
        votes.extend(parse_rating_sheet(grid, voter, songs, cfg.voting.min, cfg.voting.max, warn))
    return voters, votes
