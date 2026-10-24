import pandas as pd
import pytest

from chatting.config import Config
from chatting.models import Song
from chatting.names import People
from chatting.ratings import (
    load_votes, parse_rating_sheet, parse_rating_value, voter_name_from_path, write_rating_sheet,
)
from chatting.sources import read_grid

SONGS = [
    Song(1, "Hello", "Adele", "https://youtu.be/aaaaaaaaaaa", "Ana", 1),
    Song(2, "Yellow", "Coldplay", "https://youtu.be/bbbbbbbbbbb", "Bob", 1),
    Song(3, "Roar", "Katy Perry", "www.youtube.com/watch?v=ccccccccccc", "Cy", 1),
]


@pytest.mark.parametrize("raw,expected", [
    ("6.78", 6.78), ("6,78", 6.78), (" 7 ", 7), (7.0, 7), ("7/10", 7), ("7.5/10", 7.5), ("10!", 10),
    ("8 (lol)", 8), ("8.5 lol", 8.5), ("0", 0), ("6.789", 6.79), ("8 lol", 8),
])
def test_values(raw, expected):
    assert parse_rating_value(raw) == (expected, None)


@pytest.mark.parametrize("raw", ["", "-", "–", "—", "x", "N/A", "na", "forgot", "skip", "?", "nan", None])
def test_missing(raw):
    assert parse_rating_value(raw) == (None, None)


def test_problems():
    assert parse_rating_value("abc") == (None, "can't read 'abc'")
    assert parse_rating_value("11") == (None, "11 is outside 0–10")
    assert parse_rating_value("-1")[0] is None
    assert parse_rating_value("7/5")[0] is None
    for raw in ["2026-10-06 00:00:00", "6-Oct", "06/10/2026"]:
        value, problem = parse_rating_value(raw)
        assert value is None and "looks like a date" in problem and "retype" in problem


def test_write_sheet_and_parse_back(tmp_path):
    import openpyxl

    x, c = tmp_path / "out" / "sheet.xlsx", tmp_path / "out" / "sheet.csv"
    write_rating_sheet(SONGS, x, c)
    wb = openpyxl.load_workbook(x)
    ws = wb["Ratings"]
    assert [cell.value for cell in ws[1]] == ["#", "Song", "Link", "Rating", "Comments"]
    assert ws.freeze_panes == "A2"
    assert ws["B2"].value == "Adele - Hello"
    assert ws["C2"].hyperlink is not None
    assert ws.data_validations.dataValidation[0].sqref is not None
    assert c.read_bytes().startswith(b"\xef\xbb\xbf")

    ws["D2"], ws["E2"] = 6.78, "great"
    ws["D3"], ws["D4"] = "-", 8
    wb.save(x)
    votes = parse_rating_sheet(read_grid(str(x), tmp_path), "Zed", SONGS, 0, 10, lambda m: None)
    assert [(v.song_id, v.rating, v.comment) for v in votes] == [(1, 6.78, "great"), (2, None, ""), (3, 8, "")]
    assert all(v.voter == "Zed" for v in votes)


def _grid(rows):
    return pd.DataFrame(rows, dtype=object).fillna("")


def test_reordered_sheet():
    warnings = []
    grid = _grid([
        ["#", "Song", "Link", "Rating", "Comments"],
        [3, "Katy Perry - Roar", "", "9", "c3"],
        [1, "Adele - Hello", "", "5", ""],
        [2, "Coldplay - Yellow", "", "7", "c2"],
    ])
    votes = parse_rating_sheet(grid, "Zed", SONGS, 0, 10, warnings.append)
    assert {v.song_id: v.rating for v in votes} == {1: 5, 2: 7, 3: 9}
    assert warnings == []


def test_wrong_number_uses_text_and_warns():
    warnings = []
    grid = _grid([["#", "Song", "Rating"], [2, "Adele - Hello", "5"], [2, "Zzz Unknown", "7"]])
    votes = parse_rating_sheet(grid, "Zed", SONGS, 0, 10, warnings.append)
    assert {v.song_id: v.rating for v in votes} == {1: 5, 2: 7}
    assert any("Zed" in w for w in warnings) and len(warnings) == 2


def test_missing_id_column_and_header_offset():
    warnings = []
    grid = _grid([
        ["My ratings"],
        [],
        ["Song", "Nota", "Notes"],
        ["Adele - Hello", "5", ""],
        ["Coldplay - Yellow", "6", ""],
        ["Katy Perry - Roar", "x", "meh"],
    ])
    votes = parse_rating_sheet(grid, "Zed", SONGS, 0, 10, warnings.append)
    assert [(v.song_id, v.rating, v.comment) for v in votes] == [(1, 5, ""), (2, 6, ""), (3, None, "meh")]
    assert warnings == []


def test_missing_id_by_position():
    warnings = []
    grid = _grid([["Song", "Rating"], ["adel hello", "5"], ["???", "6"], ["Katy Perry - Roar", "7"]])
    votes = parse_rating_sheet(grid, "Zed", SONGS, 0, 10, warnings.append)
    assert {v.song_id: v.rating for v in votes} == {1: 5, 2: 6, 3: 7}
    assert len(warnings) == 1 and "position" in warnings[0]


def test_skipped_duplicate_and_bad_rating():
    warnings = []
    grid = _grid([
        ["#", "Song", "Rating"],
        [1, "Adele - Hello", "11"],
        [1, "Adele - Hello", "5"],
        [9, "Mystery", "5"],
        ["", "", ""],
    ])
    votes = parse_rating_sheet(grid, "Zed", SONGS, 0, 10, warnings.append)
    assert len(votes) == 1 and votes[0].rating is None and votes[0].raw == "11"
    text = "\n".join(warnings)
    assert "outside" in text and "twice" in text and "Mystery" in text
    assert all(w.startswith("Zed") for w in warnings)


def test_no_rating_column():
    warnings = []
    assert parse_rating_sheet(_grid([["a", "b"], ["1", "2"]]), "Zed", SONGS, 0, 10, warnings.append) == []
    assert "Zed" in warnings[0]


def test_voter_name_from_path(tmp_path):
    assert voter_name_from_path(tmp_path / "Ana (1).xlsx") == "Ana"
    assert voter_name_from_path(tmp_path / " Călin M.csv") == "Călin M"


def test_load_votes(tmp_path):
    cfg = Config(base_dir=tmp_path)
    folder = tmp_path / "data" / "ratings"
    folder.mkdir(parents=True)
    write_rating_sheet(SONGS, folder / "rating_sheet.xlsx")
    for name, r in (("ana.csv", "5"), ("Bob.csv", "6"), ("Bob (1).csv", "7"), (".hidden.csv", "1"), ("~$lock.xlsx", "1")):
        (folder / name).write_text(f"#,Song,Rating,Comments\n1,Adele - Hello,{r},hi\n", encoding="utf-8")
    (folder / "notes.txt").write_text("ignore")
    import os

    for name in ("Bob.csv", "Bob (1).csv"):
        os.utime(folder / name, (1_700_000_000, 1_700_000_000))
    warnings = []
    voters, votes = load_votes(cfg, SONGS, People({"Ana": ["ana"]}), warnings.append)
    assert voters == ["Ana", "Bob"]
    # Same time: the re-sent "(1)" copy wins.
    assert [(v.voter, v.rating) for v in votes] == [("Ana", 5), ("Bob", 7)]
    assert any("using the newest, Bob (1).csv" in w for w in warnings)

    os.utime(folder / "Bob.csv", (1_700_000_100, 1_700_000_100))
    voters, votes = load_votes(cfg, SONGS, People({"Ana": ["ana"]}), lambda m: None)
    assert [(v.voter, v.rating) for v in votes] == [("Ana", 5), ("Bob", 6)]


def test_load_votes_missing_folder(tmp_path):
    warnings = []
    assert load_votes(Config(base_dir=tmp_path), SONGS, People(), warnings.append) == ([], [])
    assert warnings


@pytest.mark.parametrize("raw", ["6 or 7", "7 8", "1:23", "4.mai", "9 (10 for the chorus)", "7-8"])
def test_rating_with_a_second_number_is_reported(raw):
    value, problem = parse_rating_value(raw)
    assert value is None and problem


@pytest.mark.parametrize("raw,expected", [("7½", 7.5), ("8 (lol)", 8), ("7/10", 7), ("6,78", 6.78),
                                          ("8. solid", 8), ("10!", 10)])
def test_rating_with_harmless_extras(raw, expected):
    assert parse_rating_value(raw) == (expected, None)


def test_generic_file_name_is_flagged(tmp_path):
    cfg = Config(base_dir=tmp_path)
    folder = tmp_path / "data" / "ratings"
    folder.mkdir(parents=True)
    (folder / "Copy of Chatting ratings.csv").write_text(
        "#,Song,Rating,Comments\n1,Adele - Hello,5,\n", encoding="utf-8")
    warnings = []
    load_votes(cfg, SONGS, People(), warnings.append)
    assert any("doesn't look like a voter's name" in w for w in warnings)
