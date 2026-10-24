import shutil

import openpyxl

from chatting.config import VotingConfig
from chatting.master import compile_results, read_master, write_master
from chatting.models import Song, Vote
from chatting.names import People

PEOPLE = People()
SONGS = [
    Song(1, "One", "A", "https://youtu.be/aaaaaaaaaaa", "Ana", 1, 45),
    Song(2, "Two", "B", "https://youtu.be/bbbbbbbbbbb", "Bob", 1, None),
    Song(3, "Three", "C", "https://youtu.be/ccccccccccc", "Cy", 1, 90),
    Song(4, "Four", "D", "https://youtu.be/ddddddddddd", "Cy", 2, None),
    Song(5, "Five", "E", "not a link", "Ana", 2, None),
]


def v(voter, sid, rating, comment=""):
    return Vote(voter, sid, rating, "" if rating is None else str(rating), comment)


VOTES = [
    v("Ana", 1, 10, "my own"), v("Ana", 2, 8), v("Ana", 3, 6), v("Ana", 4, 6),
    v("Bob", 1, 6), v("Bob", 2, 10), v("Bob", 3, 6), v("Bob", 4, 2, "bad"),
    v("Cy", 1, 8, "nice"), v("Cy", 2, 4), v("Cy", 3, 10), v("Cy", 4, 10),
]


def compile_(voting=None, votes=VOTES, voters=("Cy", "Ana", "Bob")):
    warnings = []
    res = compile_results(SONGS, [PEOPLE.canonical(x) for x in voters], votes, PEOPLE,
                          voting or VotingConfig(), warnings.append)
    return res, warnings


def by_id(res):
    return {r.song.id: r for r in res.songs}


def test_own_excluded_and_average():
    res, warnings = compile_()
    r = by_id(res)
    assert r[1].excluded == {"Ana"} and r[1].ratings["Ana"] is None
    assert r[1].average == 7 and r[1].votes == 2        # (6 + 8) / 2
    assert r[1].comments["Ana"] == "my own"             # comment kept
    assert r[2].average == 6 and r[4].average == 4      # Cy's own 10 is left out of song 4
    assert r[3].average == 6 and r[3].votes == 2
    assert any("Nobody rated song 5" in w for w in warnings)


def test_exclude_own_off():
    res, _ = compile_(VotingConfig(exclude_own=False))
    assert by_id(res)[1].average == 8 and by_id(res)[1].votes == 3


def test_ranks_ties_and_unrated_last():
    res, _ = compile_()
    # averages: 1 -> 7, 2 -> 6, 3 -> 6, 4 -> 4, 5 -> None
    assert [(r.song.id, r.rank) for r in res.songs] == [(1, 1), (2, 2), (3, 2), (4, 4), (5, 5)]


def test_unrated_share_last_rank():
    songs = SONGS + [Song(6, "Six", "F", "x", "Dee", 1)]
    res = compile_results(songs, ["Ana"], [v("Ana", 2, 5)], PEOPLE, VotingConfig(), lambda m: None)
    ranks = {r.song.id: r.rank for r in res.songs}
    assert ranks[2] == 1 and {ranks[i] for i in (1, 3, 4, 5, 6)} == {2}


def test_voter_order():
    res, _ = compile_(voters=("zed", "Cy", "Ana", "Bob"))
    assert res.voters == ["Ana", "Bob", "Cy", "zed"]
    res, _ = compile_(VotingConfig(voter_order=["bob", "Missing", "zed"]), voters=("zed", "Cy", "Ana", "Bob"))
    assert res.voters == ["Bob", "zed", "Ana", "Cy"]


def snapshot(res):
    return [(r.rank, r.song.id, r.average, r.votes, sorted(r.excluded), r.ratings, r.comments) for r in res.songs]


def test_roundtrip(tmp_path):
    res, _ = compile_()
    path = tmp_path / "out" / "master.xlsx"
    write_master(res, path, "Title")
    back = read_master(path, PEOPLE, VotingConfig(), lambda m: None)
    assert back.voters == res.voters
    assert snapshot(back) == snapshot(res)
    key = lambda s: (s.id, s.title, s.artist, s.link, s.submitter, s.start)  # noqa: E731 (slot is not stored)
    assert [key(r.song) for r in back.songs] == [key(r.song) for r in res.songs]

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Ranking", "Comments", "Submitters", "Voters"]
    ws = wb["Ranking"]
    assert [c.value for c in ws[1]][:9] == ["Rank", "Average", "Votes", "Song", "Artist", "Submitted by", "Link", "Start", "ID"]
    assert ws["B2"].number_format == "0.00"
    assert ws["H2"].value == "0:45"
    assert "own" in [c.value for c in ws[2]]
    sub = {row[0].value: row for row in wb["Submitters"].iter_rows(min_row=2)}
    assert sub["Bob"][2].value == 6 and sub["Cy"][3].value == 2 and sub["Cy"][4].value == 4


def test_backup_on_overwrite(tmp_path):
    res, _ = compile_()
    path = tmp_path / "master.xlsx"
    write_master(res, path)
    assert not (tmp_path / "master.backup.xlsx").exists()
    write_master(res, path)
    assert (tmp_path / "master.backup.xlsx").exists()


def edit(path, fn):
    wb = openpyxl.load_workbook(path)
    fn(wb["Ranking"])
    wb.save(path)


def col_of(ws, name):
    return next(c.column for c in ws[1] if c.value == name)


def row_of(ws, sid):
    return next(r[0].row for r in ws.iter_rows(min_row=2) if r[8].value == sid)


def test_edit_rating_changes_average_and_ranks(tmp_path):
    res, _ = compile_()
    path = tmp_path / "master.xlsx"
    write_master(res, path)
    # Bob's rating of song 1 -> 10; the stale Average/Rank cells must be ignored
    edit(path, lambda ws: ws.cell(row=row_of(ws, 1), column=col_of(ws, "Bob"), value=10))
    back = by_id(read_master(path, PEOPLE, VotingConfig(), lambda m: None))
    assert back[1].average == 9
    # Typing a number into an "own" cell must not count
    edit(path, lambda ws: ws.cell(row=row_of(ws, 1), column=col_of(ws, "Ana"), value=1))
    back = by_id(read_master(path, PEOPLE, VotingConfig(), lambda m: None))
    assert back[1].average == 9 and back[1].excluded == {"Ana"} and back[1].votes == 2


def test_reordered_voter_columns_respected(tmp_path):
    res, _ = compile_()
    path = tmp_path / "master.xlsx"
    write_master(res, path)

    def swap(ws):
        a, c = col_of(ws, "Ana"), col_of(ws, "Cy")
        for row in ws.iter_rows(min_row=1):
            row[a - 1].value, row[c - 1].value = row[c - 1].value, row[a - 1].value

    edit(path, swap)
    back = read_master(path, PEOPLE, VotingConfig(), lambda m: None)
    assert back.voters == ["Cy", "Bob", "Ana"] or back.voters == ["Ana", "Bob", "Cy"][::-1]
    forced = read_master(path, PEOPLE, VotingConfig(voter_order=["Bob"]), lambda m: None)
    assert forced.voters[0] == "Bob"
